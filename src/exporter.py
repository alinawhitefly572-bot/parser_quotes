import csv
from docx import Document
from pathlib import Path

def save_to_docx(data: list[dict], filepath: Path):
    mydoc = Document()
    for item in data:
        mydoc.add_paragraph(item['Текст'])
        mydoc.add_paragraph(f"— {item['Автор']}")
        mydoc.add_paragraph(f"Теги: {item['Теги']}")
        mydoc.add_paragraph("") # Пустая строка
    mydoc.save(filepath)

def save_to_csv(data: list[dict], filepath: Path):
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['Текст', 'Автор', 'Теги'])
        writer.writeheader()
        writer.writerows(data)